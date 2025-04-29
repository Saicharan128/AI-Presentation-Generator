from flask import Flask, render_template, request, send_from_directory, jsonify
from llama_cpp import Llama
import yaml
from pptx import Presentation
from pptx.util import Pt, Inches
from docx import Document
from fpdf import FPDF
import re
import os
import uuid
import requests
import io
from PIL import Image

app = Flask(__name__)
app.config['UPLOAD_FOLDER'] = 'generated_files'
app.config['PEXELS_API_KEY'] = 'THTqiDKdJPqa0wGgJYQnbSdzcpfIBZnvLuswsf6ho3JbYxnyZQGqcdax'  # Replace with your real Pexels API key

# Ensure upload directory exists
os.makedirs(app.config['UPLOAD_FOLDER'], exist_ok=True)

# Load model once
model = Llama(model_path="customweights.gguf", n_ctx=1024)

def extract_yaml_block(text):
    if "---" in text:
        yaml_parts = text.split("---", 1)
        return "---" + yaml_parts[1].split("\n---")[0]
    else:
        return "---\npresentation:\n  slides: []"

def fix_yaml_format(yaml_text):
    fixed_text = re.sub(r"^\s*[\*\d]\.", "  -", yaml_text, flags=re.MULTILINE)
    fixed_text = re.sub(r"^\s*\*", "  -", fixed_text, flags=re.MULTILINE)
    return fixed_text

def generate_yaml_from_topic(topic):
    prompt = (
        f"Create a YAML for a detailed presentation on '{topic}'.\n"
        f"The bullets MUST be complete explanatory sentences.\n"
        f"Format:\n"
        f"---\n"
        f"presentation:\n"
        f"  slides:\n"
        f"    - title: Slide 1 Title\n"
        f"      bullets:\n"
        f"        - Bullet 1 full sentence.\n"
        f"        - Bullet 2 detailed info.\n"
        f"---\n"
        f"Only output valid YAML. No extra text."
    )

    chat_messages = [
        {"role": "system", "content": "You generate YAML content for informative presentations. Never output questions, only clear explanations."},
        {"role": "user", "content": prompt}
    ]

    response = model.create_chat_completion(
        messages=chat_messages,
        max_tokens=700,
        temperature=0.3,
        top_p=0.9,
        stop=["</s>"]
    )

    raw_output = response['choices'][0]['message']['content'].strip()
    yaml_only = extract_yaml_block(raw_output)
    yaml_fixed = fix_yaml_format(yaml_only)

    return yaml_fixed

def search_pexels_image(query):
    try:
        headers = {'Authorization': app.config['PEXELS_API_KEY']}
        url = f"https://api.pexels.com/v1/search?query={query}&per_page=1"
        response = requests.get(url, headers=headers)
        if response.status_code == 200:
            data = response.json()
            if data.get('photos'):
                return data['photos'][0]['src']['large']
        return None
    except Exception as e:
        print(f"Pexels search error: {e}")
        return None

def download_image(url):
    try:
        response = requests.get(url)
        return response.content if response.status_code == 200 else None
    except Exception as e:
        print(f"Image download error: {e}")
        return None

def determine_optimal_font_size(slides_data, presentation):
    font_sizes = [18, 16, 14, 12, 11, 10]
    content_width = presentation.slide_width * 0.85
    content_height = presentation.slide_height * 0.7

    max_lines = 0
    for slide_info in slides_data:
        bullets = slide_info.get('bullets', [])
        estimated_lines = 0
        for bullet in bullets:
            chars_per_line = int(content_width / Pt(12).inches * 2)
            bullet_lines = max(1, len(bullet) / chars_per_line)
            estimated_lines += bullet_lines + 0.5
        max_lines = max(max_lines, estimated_lines)

    for size in font_sizes:
        line_height = Pt(size).inches * 1.2
        max_possible = content_height / line_height
        if max_lines <= max_possible:
            return size
    return 10

import random

def fetch_consistent_background_image(count=5):
    search_terms = [
        "white abstract background",
        "light texture background",
        "minimal white backdrop",
        "soft clean gradient",
        "light pastel abstract",
        "white particle background",
        "abstract light particles",
        "minimal tech particles",
        "digital wave particles"
    ]
    
    random.shuffle(search_terms)
    images = []
    
    for term in search_terms:
        img_url = search_pexels_image(term)
        if img_url:
            img_data = download_image(img_url)
            if img_data and is_bright_image(img_data):
                images.append(img_data)
        if len(images) >= count:
            break

    return images[0] if images else None

def is_bright_image(img_data, threshold=180):
    try:
        img = Image.open(io.BytesIO(img_data)).convert('L')  # grayscale
        stat = img.resize((50, 50)).getdata()
        avg = sum(stat) / len(stat)
        return avg > threshold
    except Exception as e:
        print(f"Brightness check error: {e}")
        return False

def create_content_slides(prs, layout, title, bullets, font_size, background_img_data=None):
    font_pt = Pt(font_size)
    
    content_height = prs.slide_height * 0.6
    content_top = Inches(1.5)
    content_left = Inches(1)
    content_width = prs.slide_width - Inches(2)

    line_height = font_pt.pt * 1.3
    max_lines = int(content_height / Pt(line_height))

    current, remaining = [], bullets[:]
    while remaining:
        current, lines = [], 0
        while remaining and lines < max_lines:
            bullet = remaining[0]
            chars_per_line = int(80 * (14 / font_size))
            est_lines = max(1, len(bullet) / chars_per_line) + 0.5
            if lines + est_lines <= max_lines:
                current.append(remaining.pop(0))
                lines += est_lines
            else:
                break

        slide = prs.slides.add_slide(prs.slide_layouts[6])  # Use blank layout

        # Add background image (behind everything)
        if background_img_data:
            img_stream = io.BytesIO(background_img_data)
            slide.shapes.add_picture(img_stream, 0, 0, width=prs.slide_width, height=prs.slide_height)

        # Add title on top
        title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.5), prs.slide_width - Inches(1.6), Inches(1))
        title_frame = title_box.text_frame
        title_frame.text = f"{title} (continued)" if remaining else title
        title_frame.paragraphs[0].font.size = Pt(font_size + 2)
        title_frame.paragraphs[0].font.bold = True

        # Add bullet points
        content_box = slide.shapes.add_textbox(content_left, content_top, content_width, content_height)
        tf = content_box.text_frame
        tf.word_wrap = True
        for i, bullet in enumerate(current):
            p = tf.add_paragraph()
            p.text = bullet
            p.level = 0
            p.font.size = font_pt

def add_image_slide(prs, layout, title, image_url, background_img_data=None):
    img_data = download_image(image_url)
    if img_data:
        slide = prs.slides.add_slide(layout)

        # Add background image
        if background_img_data:
            bg_stream = io.BytesIO(background_img_data)
            slide.shapes.add_picture(bg_stream, 0, 0, width=prs.slide_width, height=prs.slide_height)

        # Add main image
        img = Image.open(io.BytesIO(img_data))
        width, height = img.size
        max_w, max_h = Inches(10), Inches(5.5)
        scale = min(max_w / width, max_h / height)
        new_w, new_h = width * scale, height * scale
        left = (prs.slide_width - new_w) / 2
        top = (prs.slide_height - new_h + Inches(1)) / 2
        slide.shapes.title.text = f"{title} - Visual"
        slide.shapes.add_picture(io.BytesIO(img_data), left, top, width=new_w, height=new_h)


def create_ppt_from_yaml(yaml_content, output_path, topic):
    try:
        data = yaml.safe_load(yaml_content)
        slides_data = data.get('presentation', {}).get('slides', [])
        if not slides_data:
            return {"success": False, "error": "No slides found"}

        prs = Presentation()
        layout = prs.slide_layouts[1]
        img_layout = prs.slide_layouts[5]
        font_size = determine_optimal_font_size(slides_data, prs)

        # Fetch consistent background image
        background_img_data = fetch_consistent_background_image()

        for slide in slides_data:
            title = slide.get('title', 'Untitled')
            bullets = slide.get('bullets', [])

            # Fetch image for the slide
            img_url = search_pexels_image(f"{topic} {title}")

            create_content_slides(prs, layout, title, bullets, font_size, background_img_data=background_img_data)

            # Optional: also add separate image-only slide
            if img_url:
                add_image_slide(prs, img_layout, title, img_url, background_img_data=background_img_data)

        prs.save(output_path)
        return {"success": True, "font_size": font_size}
    except Exception as e:
        return {"success": False, "error": str(e)}

from pptx.dml.color import RGBColor

def add_overlay(slide, left, top, width, height, color=(255, 255, 255), transparency=0.3):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    fill = shape.fill
    fill.solid()
    fill.fore_color.rgb = RGBColor(*color)
    fill.transparency = transparency
    shape.line.fill.background()  # Remove border

def create_docx_from_yaml(yaml_content, output_path):
    try:
        data = yaml.safe_load(yaml_content)
        slides = data.get('presentation', {}).get('slides', [])
        if not slides:
            return {"success": False, "error": "No slides found"}
        doc = Document()
        doc.add_heading('Presentation', 0)
        for slide in slides:
            doc.add_heading(slide.get('title', 'Untitled'), level=1)
            for bullet in slide.get('bullets', []):
                doc.add_paragraph(bullet, style='ListBullet')
        doc.save(output_path)
        return {"success": True}
    except Exception as e:
        return {"success": False, "error": str(e)}

def create_pdf_from_yaml(yaml_content, output_path):
    try:
        data = yaml.safe_load(yaml_content)
        slides = data.get('presentation', {}).get('slides', [])
        if not slides:
            return {"success": False, "error": "No slides found"}
        pdf = FPDF()
        pdf.add_page()
        pdf.set_font("Arial", size=14)
        pdf.cell(0, 10, txt="Presentation", ln=True, align='C')
        pdf.ln(10)

        for slide in slides:
            pdf.set_font("Arial", 'B', 12)
            pdf.cell(0, 10, txt=slide.get('title', 'Untitled Slide'), ln=True)
            pdf.set_font("Arial", size=11)
            for bullet in slide.get('bullets', []):
                pdf.set_x(20)
                pdf.multi_cell(0, 8, txt=f"- {bullet}")
                pdf.ln(2)
            pdf.ln(5)

        pdf.output(output_path)
        return {"success": True}
    except Exception as e:
        return {"success": False, "error": str(e)}


@app.route('/')
def index():
    return render_template('index.html')

@app.route('/generate', methods=['POST'])
def generate():
    topic = request.form.get('topic')
    file_type = request.form.get('file_type')
    
    if not topic or not file_type:
        return jsonify({"success": False, "error": "Missing topic or file type"})
    
    # Generate a unique ID for this file
    file_id = str(uuid.uuid4())
    file_name = f"{topic.replace(' ', '_')}_{file_id}"
    
    # Generate YAML content
    yaml_content = generate_yaml_from_topic(topic)
    yaml_content_preview = None
    
    try:
        # Parse YAML to get preview content
        preview_data = yaml.safe_load(yaml_content)
        yaml_content_preview = preview_data
    except Exception as e:
        return jsonify({"success": False, "error": f"YAML generation error: {str(e)}"})
    
    # Process according to file type
    output_path = os.path.join(app.config['UPLOAD_FOLDER'], file_name)
    result = {"success": False, "error": "Invalid file type"}
    
    if file_type == "pptx":
        full_path = f"{output_path}.pptx"
        result = create_ppt_from_yaml(yaml_content, full_path, topic)
        if result["success"]:
            result["file_url"] = f"/download/{file_name}.pptx"
            result["preview"] = yaml_content_preview
    
    elif file_type == "docx":
        full_path = f"{output_path}.docx"
        result = create_docx_from_yaml(yaml_content, full_path)
        if result["success"]:
            result["file_url"] = f"/download/{file_name}.docx"
            result["preview"] = yaml_content_preview
    
    elif file_type == "pdf":
        full_path = f"{output_path}.pdf"
        result = create_pdf_from_yaml(yaml_content, full_path)
        if result["success"]:
            result["file_url"] = f"/download/{file_name}.pdf"
            result["preview"] = yaml_content_preview
    
    return jsonify(result)

@app.route('/download/<filename>')
def download_file(filename):
    return send_from_directory(app.config['UPLOAD_FOLDER'], filename, as_attachment=True)

# Create templates directory if it doesn't exist
if not os.path.exists('templates'):
    os.makedirs('templates')

# Write index.html template
with open('templates/index.html', 'w', encoding='utf-8') as f:
    f.write('''
<!DOCTYPE html>
<html lang="en">
<head>
    <meta charset="UTF-8">
    <meta name="viewport" content="width=device-width, initial-scale=1.0">
    <title>AI Presentation Generator</title>
    <script src="https://cdn.tailwindcss.com"></script>
    <script src="https://unpkg.com/htmx.org@1.9.2"></script>
    <style>
        /* Futuristic Loader Animation */
        .loader {
            border: 4px solid rgba(255, 255, 255, 0.2);
            border-top: 4px solid #a78bfa;
            border-radius: 50%;
            width: 60px;
            height: 60px;
            animation: spinner 1.2s cubic-bezier(0.5, 0, 0.5, 1) infinite;
            position: relative;
            box-shadow: 0 0 20px rgba(167, 139, 250, 0.5);
        }
        .loader::after {
            content: '';
            position: absolute;
            top: -10px;
            left: -10px;
            right: -10px;
            bottom: -10px;
            border-radius: 50%;
            border: 2px solid transparent;
            border-top-color: #a78bfa;
            animation: spinner 2s ease-in-out infinite;
        }
        @keyframes spinner {
            0% { transform: rotate(0deg); }
            100% { transform: rotate(360deg); }
        }

        /* Gradient Animation for Background */
        .bg-animated {
            background: linear-gradient(45deg, #4f46e5, #7c3aed, #db2777, #3b82f6);
            background-size: 400%;
            animation: gradientShift 15s ease infinite;
        }
        @keyframes gradientShift {
            0% { background-position: 0% 50%; }
            50% { background-position: 100% 50%; }
            100% { background-position: 0% 50%; }
        }

        /* Floating Card Animation */
        .card-float {
            animation: float 6s ease-in-out infinite;
        }
        @keyframes float {
            0% { transform: translateY(0px); }
            50% { transform: translateY(-10px); }
            100% { transform: translateY(0px); }
        }

        /* Glow Effect on Hover */
        .glow-hover {
            transition: all 0.3s ease;
        }
        .glow-hover:hover {
            box-shadow: 0 0 20px rgba(167, 139, 250, 0.7), 0 0 40px rgba(167, 139, 250, 0.3);
            transform: scale(1.02);
        }

        /* Input Focus Animation */
        input:focus {
            box-shadow: 0 0 15px rgba(167, 139, 250, 0.5);
            transform: scale(1.01);
            transition: all 0.3s ease;
        }

        /* Button Pulse Animation */
        .btn-pulse {
            position: relative;
            overflow: hidden;
        }
        .btn-pulse::after {
            content: '';
            position: absolute;
            top: 50%;
            left: 50%;
            width: 0;
            height: 0;
            background: rgba(255, 255, 255, 0.3);
            border-radius: 50%;
            transform: translate(-50%, -50%);
            transition: width 0.6s ease, height 0.6s ease;
        }
        .btn-pulse:hover::after {
            width: 300px;
            height: 300px;
        }

        /* Radio Button Animation */
        .radio-label {
            transition: all 0.3s ease;
            position: relative;
            overflow: hidden;
        }
        .radio-label::before {
            content: '';
            position: absolute;
            top: 0;
            left: -100%;
            width: 100%;
            height: 100%;
            background: linear-gradient(90deg, transparent, rgba(167, 139, 250, 0.2), transparent);
            transition: all 0.5s ease;
        }
        .radio-label:hover::before {
            left: 100%;
        }
        .peer-checked\/radio-label {
            transform: scale(1.05);
            box-shadow: 0 0 15px rgba(167, 139, 250, 0.4);
        }

        /* Fade-in Animation for Containers */
        .fade-in {
            opacity: 0;
            transform: translateY(20px);
            animation: fadeIn 0.8s ease-out forwards;
        }
        @keyframes fadeIn {
            to {
                opacity: 1;
                transform: translateY(0);
            }
        }

        /* Custom Scrollbar */
        .custom-scrollbar::-webkit-scrollbar {
            width: 8px;
        }
        .custom-scrollbar::-webkit-scrollbar-track {
            background: rgba(255, 255, 255, 0.1);
            border-radius: 4px;
        }
        .custom-scrollbar::-webkit-scrollbar-thumb {
            background: #a78bfa;
            border-radius: 4px;
            box-shadow: 0 0 10px rgba(167, 139, 250, 0.5);
        }

        /* Particle Canvas Styling */
        #particleCanvas {
            position: fixed;
            top: 0;
            left: 0;
            width: 100%;
            height: 100%;
            z-index: -1;
            pointer-events: none;
        }
    </style>
</head>
<body class="bg-animated min-h-screen">
    <canvas id="particleCanvas"></canvas>
    <div class="container mx-auto px-4 py-8">
        <header class="text-center mb-12 fade-in">
            <h1 class="text-5xl font-extrabold text-white mb-3 tracking-tight drop-shadow-lg">
                AI Presentation Generator
            </h1>
            <p class="text-xl text-indigo-200 font-light">Create stunning presentations with AI magic</p>
        </header>
        
        <div class="max-w-3xl mx-auto bg-white/95 backdrop-blur-lg rounded-2xl shadow-2xl p-10 card-float glow-hover">
            <form id="generationForm" class="space-y-8">
                <div class="fade-in" style="animation-delay: 0.2s;">
                    <label for="topic" class="block text-sm font-medium text-gray-800 mb-2">Presentation Topic</label>
                    <input type="text" id="topic" name="topic" required
                        class="w-full px-5 py-4 rounded-xl border border-gray-200 bg-gray-50/50 focus:ring-4 focus:ring-indigo-300 focus:border-indigo-600 text-gray-900 placeholder-gray-400 transition-all duration-300">
                </div>
                
                <div class="fade-in" style="animation-delay: 0.4s;">
                    <label class="block text-sm font-medium text-gray-800 mb-2">Output Format</label>
                    <div class="grid grid-cols-3 gap-6">
                        <div>
                            <input type="radio" id="pptx" name="file_type" value="pptx" checked class="hidden peer">
                            <label for="pptx" class="radio-label block cursor-pointer text-center p-4 border border-gray-200 rounded-xl peer-checked:bg-indigo-50 peer-checked:border-indigo-600 peer-checked/radio-label hover:bg-gray-50 transition-all duration-300">
                                <span class="block text-xl font-semibold text-gray-900">PowerPoint</span>
                                <span class="text-sm text-gray-500">.pptx</span>
                                <span class="block text-xs text-indigo-600 mt-2 font-medium">Includes Pexels images</span>
                            </label>
                        </div>
                        
                        <div>
                            <input type="radio" id="docx" name="file_type" value="docx" class="hidden peer">
                            <label for="docx" class="radio-label block cursor-pointer text-center p-4 border border-gray-200 rounded-xl peer-checked:bg-indigo-50 peer-checked:border-indigo-600 peer-checked/radio-label hover:bg-gray-50 transition-all duration-300">
                                <span class="block text-xl font-semibold text-gray-900">Word</span>
                                <span class="text-sm text-gray-500">.docx</span>
                            </label>
                        </div>
                        
                        <div>
                            <input type="radio" id="pdf" name="file_type" value="pdf" class="hidden peer">
                            <label for="pdf" class="radio-label block cursor-pointer text-center p-4 border border-gray-200 rounded-xl peer-checked:bg-indigo-50 peer-checked:border-indigo-600 peer-checked/radio-label hover:bg-gray-50 transition-all duration-300">
                                <span class="block text-xl font-semibold text-gray-900">PDF</span>
                                <span class="text-sm text-gray-500">.pdf</span>
                            </label>
                        </div>
                    </div>
                </div>
                
                <div class="fade-in" style="animation-delay: 0.6s;">
                    <button type="submit" id="generateBtn" 
                        class="w-full bg-gradient-to-r from-indigo-600 to-purple-600 hover:from-indigo-700 hover:to-purple-700 text-white font-semibold py-4 px-6 rounded-xl btn-pulse transition-all duration-300 shadow-lg">
                        Generate Presentation
                    </button>
                </div>
            </form>
            
            <div id="loadingIndicator" class="hidden mt-10 text-center fade-in">
                <div class="loader mx-auto"></div>
                <p class="mt-4 text-gray-700 font-medium">Crafting your presentation...</p>
                <p class="text-sm text-gray-500 mt-2">Please wait, this may take a moment</p>
            </div>
            
            <div id="resultContainer" class="hidden mt-10 fade-in">
                <div class="bg-green-50/80 backdrop-blur-sm border border-green-200 rounded-xl p-6 mb-6">
                    <div class="flex items-center">
                        <svg class="w-6 h-6 text-green-500 mr-3" fill="currentColor" viewBox="0 0 20 20">
                            <path fill-rule="evenodd" d="M10 18a8 8 0 100-16 8 8 0 000 16zm3.707-9.293a1 1 0 00-1.414-1.414L9 10.586 7.707 9.293a1 1 0 00-1.414 1.414l2 2a1 1 0 001.414 0l4-4z" clip-rule="evenodd"/>
                        </svg>
                        <p class="text-green-700 font-semibold text-lg">Your presentation is ready!</p>
                    </div>
                </div>
                
                <a id="downloadLink" href="#" 
                   class="block w-full bg-gradient-to-r from-indigo-600 to-purple-600 hover:from-indigo-700 hover:to-purple-700 text-white text-center font-semibold py-4 px-6 rounded-xl btn-pulse mb-6 transition-all duration-300 shadow-lg">
                    Download
                </a>
                
                <div>
                    <h3 class="text-xl font-semibold text-gray-900 mb-4">Content Preview</h3>
                    <div id="previewContent" class="bg-gray-50/80 backdrop-blur-sm rounded-xl p-6 max-h-96 overflow-y-auto border custom-scrollbar">
                        <!-- Preview content will be populated here -->
                    </div>
                    <p class="text-sm text-gray-500 mt-3">Note: PowerPoint presentations include relevant images from Pexels after each content slide.</p>
                </div>
            </div>
            
            <div id="errorContainer" class="hidden mt-10 fade-in">
                <div class="bg-red-50/80 backdrop-blur-sm border border-red-200 rounded-xl p-6">
                    <div class="flex items-center">
                        <svg class="w-6 h-6 text-red-500 mr-3" fill="currentColor" viewBox="0 0 20 20">
                            <path fill-rule="evenodd" d="M10 18a8 8 0 100-16 8 8 0 000 16zM8.707 7.293a1 1 0 00-1.414 1.414L8.586 10l-1.293 1.293a1 1 0 101.414 1.414L10 11.414l1.293 1.293a1 1 0 001.414-1.414L11.414 10l1.293-1.293a1 1 0 00-1.414-1.414L10 8.586 8.707 7.293z" clip-rule="evenodd"/>
                        </svg>
                        <p class="text-red-700 font-semibold text-lg" id="errorMessage">Something went wrong</p>
                    </div>
                </div>
            </div>
        </div>
    </div>
    
    <script>
        // Particle Animation
        const canvas = document.getElementById('particleCanvas');
        const ctx = canvas.getContext('2d');
        canvas.width = window.innerWidth;
        canvas.height = window.innerHeight;

        const particles = [];
        const particleCount = 100;

        class Particle {
            constructor() {
                this.x = Math.random() * canvas.width;
                this.y = Math.random() * canvas.height;
                this.size = Math.random() * 3 + 1;
                this.speedX = Math.random() * 0.5 - 0.25;
                this.speedY = Math.random() * 0.5 - 0.25;
                this.opacity = Math.random() * 0.5 + 0.3;
            }

            update() {
                this.x += this.speedX;
                this.y += this.speedY;
                this.opacity = Math.sin(Date.now() * 0.001 + this.x) * 0.3 + 0.5;

                if (this.x < 0 || this.x > canvas.width) this.speedX *= -1;
                if (this.y < 0 || this.y > canvas.height) this.speedY *= -1;
            }

            draw() {
                ctx.beginPath();
                ctx.arc(this.x, this.y, this.size, 0, Math.PI * 2);
                ctx.fillStyle = `rgba(167, 139, 250, ${this.opacity})`;
                ctx.shadowBlur = 10;
                ctx.shadowColor = 'rgba(167, 139, 250, 0.8)';
                ctx.fill();
            }
        }

        function initParticles() {
            for (let i = 0; i < particleCount; i++) {
                particles.push(new Particle());
            }
        }

        function animateParticles() {
            ctx.clearRect(0, 0, canvas.width, canvas.height);
            particles.forEach(particle => {
                particle.update();
                particle.draw();
            });
            requestAnimationFrame(animateParticles);
        }

        initParticles();
        animateParticles();

        window.addEventListener('resize', () => {
            canvas.width = window.innerWidth;
            canvas.height = window.innerHeight;
        });

        // Form Submission Logic
        document.getElementById('generationForm').addEventListener('submit', function(e) {
            e.preventDefault();
            
            // Show loading indicator
            document.getElementById('loadingIndicator').classList.remove('hidden');
            document.getElementById('resultContainer').classList.add('hidden');
            document.getElementById('errorContainer').classList.add('hidden');
            document.getElementById('generateBtn').disabled = true;
            document.getElementById('generateBtn').classList.add('opacity-70');
            
            // Get form data
            const formData = new FormData(this);
            
            // Send request
            fetch('/generate', {
                method: 'POST',
                body: formData
            })
            .then(response => response.json())
            .then(data => {
                // Hide loading indicator
                document.getElementById('loadingIndicator').classList.add('hidden');
                document.getElementById('generateBtn').disabled = false;
                document.getElementById('generateBtn').classList.remove('opacity-70');
                
                if (data.success) {
                    // Show result
                    document.getElementById('resultContainer').classList.remove('hidden');
                    document.getElementById('downloadLink').href = data.file_url;
                    
                    // Show preview
                    const previewContent = document.getElementById('previewContent');
                    previewContent.innerHTML = '';
                    
                    if (data.preview && data.preview.presentation && data.preview.presentation.slides) {
                        const slides = data.preview.presentation.slides;
                        slides.forEach((slide, index) => {
                            const slideElement = document.createElement('div');
                            if (index > 0) {
                                slideElement.classList.add('mt-4', 'pt-4', 'border-t', 'border-gray-200');
                            }
                            
                            const titleElement = document.createElement('h4');
                            titleElement.textContent = slide.title || 'Untitled Slide';
                            titleElement.classList.add('font-semibold', 'text-gray-900', 'mb-3');
                            slideElement.appendChild(titleElement);
                            
                            if (slide.bullets && slide.bullets.length) {
                                const bulletList = document.createElement('ul');
                                bulletList.classList.add('list-disc', 'pl-6', 'space-y-2');
                                
                                slide.bullets.forEach(bullet => {
                                    const bulletItem = document.createElement('li');
                                    bulletItem.textContent = bullet;
                                    bulletItem.classList.add('text-gray-700', 'textγη-sm');
                                    bulletList.appendChild(bulletItem);
                                });
                                
                                slideElement.appendChild(bulletList);
                            }
                            
                            previewContent.appendChild(slideElement);
                            
                            // Add note about image slide in PowerPoint only
                            if (formData.get('file_type') === 'pptx') {
                                const imageNote = document.createElement('p');
                                imageNote.textContent = "Will include Pexels image slide after this content";
                                imageNote.classList.add('text-xs', 'text-indigo-600', 'mt-2', 'italic');
                                slideElement.appendChild(imageNote);
                            }
                        });
                    } else {
                        previewContent.textContent = "No preview available";
                    }
                } else {
                    // Show error
                    document.getElementById('errorContainer').classList.remove('hidden');
                    document.getElementById('errorMessage').textContent = data.error || "Something went wrong";
                }
            })
            .catch(error => {
                // Hide loading indicator and show error
                document.getElementById('loadingIndicator').classList.add('hidden');
                document.getElementById('errorContainer').classList.remove('hidden');
                document.getElementById('errorMessage').textContent = "Network error occurred";
                document.getElementById('generateBtn').disabled = false;
                document.getElementById('generateBtn').classList.remove('opacity-70');
                console.error('Error:', error);
            });
        });
    </script>
</body>
</html>
    ''')

if __name__ == "__main__":
    app.run(debug=True)